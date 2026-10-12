"""Reports and exports: CSV (with a spreadsheet-injection guard: a text cell that starts with = + - @ is stored as text) and a multi-sheet Excel pack. Figures come from the same services the screens use."""
import csv
import io
from datetime import date
from decimal import Decimal
from typing import Optional

from accfino.modules.investments.services import core, data, integration, portfolio, property_svc, wealth


def _cell(v):
    if isinstance(v, (Decimal, int, float)) or v is None:
        return "" if v is None else v
    s = str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def to_csv(rows, columns) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[1] for c in columns])
    for r in rows:
        w.writerow([_cell(r.get(c[0])) for c in columns])
    return buf.getvalue()


HOLD = [("account", "Account"), ("code", "Code"), ("name", "Name"), ("kind", "Type"), ("asset_class", "Asset class"), ("qty", "Units"), ("cost_base", "Cost base"), ("unit_cost", "Cost per unit"), ("price", "Price"), ("price_date", "Price date"),
        ("value", "Market value"), ("unrealised", "Unrealised gain/(loss)"), ("unrealised_pct", "Unrealised %"), ("income_12m", "Income last 12 months"), ("weight", "Weight")]
DISP = [("dispose_date", "Disposal date"), ("code", "Code"), ("name", "Name"), ("kind", "Type"), ("account", "Account"), ("acquire_date", "Acquired"), ("qty", "Units"), ("proceeds", "Proceeds"), ("selling_costs", "Selling costs"), ("cost_base", "Cost base"),
        ("gain", "Gain/(loss)"), ("discount_eligible", "Held >12 months (discount eligible)"), ("days_held", "Days held"), ("method", "Parcel method"), ("fy", "Income year")]
INC = [("date", "Date"), ("fy", "Income year"), ("code", "Code"), ("name", "Name"), ("type", "Type"), ("cash", "Cash"), ("franked", "Franked"), ("unfranked", "Unfranked"), ("franking_credits", "Franking credits"), ("foreign_income", "Foreign income"),
       ("cg_discount_part", "Capital gain (discountable)"), ("cg_other_part", "Capital gain (other)"), ("tax_deferred", "Tax deferred"), ("withholding", "Tax withheld"), ("value_in_kind", "Income in kind (staking/airdrop)")]
TXN = [("trade_date", "Date"), ("account", "Account"), ("code", "Code"), ("security", "Name"), ("txn_type", "Type"), ("quantity", "Units"), ("price", "Price"), ("amount", "Amount"), ("fees", "Fees"), ("franked", "Franked"), ("unfranked", "Unfranked"),
       ("franking_credits", "Franking credits"), ("tax_deferred", "Tax deferred"), ("cost_base_adj", "AMIT cost base net"), ("reference", "Reference"), ("source", "Source"), ("voided", "Voided"), ("linked_journal_id", "Ledger journal")]
PARCEL = [("code", "Code"), ("account", "Account"), ("acquired", "Acquired"), ("qty", "Units"), ("cost_base", "Cost base"), ("unit_cost", "Cost per unit"), ("origin", "Origin"), ("discount_eligible", "Over 12 months")]


def csv_for(db, ctx, name: str, fy: str = "", kind: str = "", as_at: Optional[date] = None) -> tuple:
    as_at = as_at or date.today()
    if name == "holdings":
        return "holdings.csv", to_csv(core.plain(portfolio.holdings(db, ctx, as_at, kind)["rows"]), HOLD)
    if name == "realised":
        return "realised_gains.csv", to_csv(core.plain(portfolio.disposals(db, ctx, fy, kind)["rows"]), DISP)
    if name == "income":
        return "income.csv", to_csv(core.plain(portfolio.income(db, ctx, fy)["rows"]), INC)
    if name == "transactions":
        return "transactions.csv", to_csv(data.list_txns(db, ctx, kind=kind, include_voided=True), TXN)
    if name == "parcels":
        secs = {s.id: s for s in db.query(data.M.InvSecurity).filter_by(org_id=ctx.org.id)}
        accs = {a.id: a for a in db.query(data.M.InvAccount).filter_by(org_id=ctx.org.id)}
        res = portfolio.replay(db, ctx, as_at)
        rows = [dict(code=secs[l.security_id].code, account=accs[l.account_id].name, acquired=l.acquired.isoformat(), qty=l.qty, cost_base=l.cost, unit_cost=l.unit_cost, origin=l.origin, discount_eligible=as_at > portfolio.add_months(l.acquired, 12))
                for l in sorted(res.lots, key=lambda l: (secs[l.security_id].code, l.acquired))]
        return "parcels_cost_base.csv", to_csv(core.plain(rows), PARCEL)
    if name == "networth":
        n = core.plain(wealth.networth(db, ctx, as_at))
        rows = [dict(side="Asset", **r) for r in n["assets"]] + [dict(side="Liability", **r) for r in n["liabilities"]]
        return "net_worth.csv", to_csv(rows, [("side", "Side"), ("group", "Group"), ("name", "Item"), ("value", "Value"), ("source", "Source")])
    raise core.NotFound("Report")


def xlsx_pack(db, ctx, fy: str = "", as_at: Optional[date] = None) -> bytes:
    from openpyxl import Workbook
    as_at = as_at or date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    n = wealth.networth(db, ctx, as_at)
    h = portfolio.holdings(db, ctx, as_at)
    perf = portfolio.performance(db, ctx, as_at)
    for r in (("AccFino Investments & Wealth pack",), ("As at", as_at.isoformat()), ("Income year filter", fy or "all"), (), ("Investments market value", float(h["total_value"])), ("Investments cost base", float(h["total_cost"])),
              ("Net worth", float(n["net_worth"])), ("Total assets", float(n["total_assets"])), ("Total liabilities", float(n["total_liabilities"])), ("XIRR since inception", perf["xirr_since_inception"]), (),
              ("Prepared from your AccFino records. Realised gains here are an estimate: the authoritative CGT calculation is in Taxation > CGT.",)):
        ws.append(list(r))
    for title, cols, rows in (("Holdings", HOLD, portfolio.holdings(db, ctx, as_at)["rows"]), ("Realised gains", DISP, portfolio.disposals(db, ctx, fy)["rows"]), ("Income", INC, portfolio.income(db, ctx, fy)["rows"]), ("Transactions", TXN, data.list_txns(db, ctx))):
        s = wb.create_sheet(title)
        s.append([c[1] for c in cols])
        for r in core.plain(rows):
            s.append([_cell(r.get(c[0])) if not isinstance(r.get(c[0]), str) or not r.get(c[0]).replace(".", "", 1).replace("-", "", 1).isdigit() else float(r.get(c[0])) for c in cols])
    s = wb.create_sheet("Net worth")
    s.append(["Side", "Group", "Item", "Value"])
    for side, rows in (("Asset", n["assets"]), ("Liability", n["liabilities"])):
        for r in rows:
            s.append([side, r["group"], _cell(r["name"]), float(r["value"])])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
